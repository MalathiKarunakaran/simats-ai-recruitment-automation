"""Selection batches (docs/specs/selection-batches.md), backend.

NOTE: like every router test in this suite these run inside one SAVEPOINT the
TestClient shares, so a missing `db.commit()` in the router would not fail
here. The commits were verified by hand against a live server (see the
hand-over note), not by this file.
"""

import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.models.audit_log import AuditLog
from app.models.enums import GLOBAL_SCOPE_ROLES, TalentPoolSourceTabEnum, UserRoleEnum
from app.models.selection_batch import SelectionBatch, SelectionBatchMember
from app.models.talent_pool_entry import TalentPoolEntry

from tests.conftest import auth_headers

URL = "/api/v1/selection-batches"
ALLOWED = (UserRoleEnum.SUPER_ADMIN, UserRoleEnum.HR_ADMIN, UserRoleEnum.RECRUITMENT_COORDINATOR)


def _entry(db, **kw) -> TalentPoolEntry:
    defaults = {
        "source_tab": TalentPoolSourceTabEnum.TEACHING_FINAL,
        "full_name": f"Candidate {uuid.uuid4().hex[:6]}",
        "sheet_status": "Selected",
        "is_active": True,
        "raw": {},
    }
    defaults.update(kw)
    e = TalentPoolEntry(**defaults)
    db.add(e)
    db.flush()
    return e


@pytest.fixture()
def admin(user_factory, client):
    return auth_headers(client, user_factory(UserRoleEnum.SUPER_ADMIN))


def _create(client, headers, entries, **extra):
    body = {"members": [{"talent_pool_entry_id": str(e.id)} for e in entries], **extra}
    return client.post(URL, json=body, headers=headers)


def test_roles_are_global_scope():
    assert set(ALLOWED) <= set(GLOBAL_SCOPE_ROLES)


def test_create_snapshot_and_fields(client, db_session, admin, campus_factory):
    campus = campus_factory("SSE")
    final = _entry(db_session, full_name="Dr Final", dept_group="CSE", dept_raw="Comp Sci", eligibility_raw="With Ph.D",
                   remarks="Strong panel", applied_post="IGNORED")  # fmt: skip
    nts = _entry(db_session, source_tab=TalentPoolSourceTabEnum.NTS, full_name="Nts Person", dept_raw="Library",
                 dept_group="Other", applied_post="Lab Assistant")  # fmt: skip
    body = {
        "members": [
            {"talent_pool_entry_id": str(final.id), "designation": "Professor", "proposed_salary": "90,000"},
            {"talent_pool_entry_id": str(nts.id)},
        ],
        "campus_id": str(campus.id),
        "note": "First list",
    }
    r = client.post(URL, json=body, headers=admin)
    assert r.status_code == 201, r.text
    d = r.json()
    assert d["batch_no"] == f"SB-{date.today().year}-001"
    assert d["status"] == "SENT_TO_UNIVERSITY" and d["campus_code"] == "SSE" and d["member_count"] == 2
    f, n = d["members"]
    assert (f["full_name"], f["designation"], f["department"], f["category"]) == ("Dr Final", "Professor", "CSE", "TEACHING")
    assert (f["qualification"], f["panel_remarks"], f["proposed_salary"], f["experience"]) == ("With Ph.D", "Strong panel", "90,000", None)
    assert (n["designation"], n["department"], n["category"]) == ("Lab Assistant", "Library", "NON_TEACHING")

    # Final with no typed designation has none (applied_post is NTS-only)
    r2 = _create(client, admin, [_entry(db_session, applied_post="X")])
    assert r2.json()["members"][0]["designation"] is None


def test_snapshot_is_frozen_when_pool_row_changes(client, db_session, admin):
    e = _entry(db_session, full_name="Before", dept_group="ECE", remarks="old remark")
    batch_id = _create(client, admin, [e]).json()["id"]
    e.full_name, e.dept_group, e.remarks, e.sheet_status = "After", "CSE", "new remark", "Rejected"
    db_session.flush()
    m = client.get(f"{URL}/{batch_id}", headers=admin).json()["members"][0]
    assert (m["full_name"], m["department"], m["panel_remarks"]) == ("Before", "ECE", "old remark")


@pytest.mark.parametrize("kw", [{"sheet_status": "Pending"}, {"sheet_status": "Rejected"}, {"is_active": False}])
def test_refuses_non_selected_or_inactive(client, db_session, admin, kw):
    e = _entry(db_session, full_name="Nope Person", **kw)
    r = _create(client, admin, [e])
    assert r.status_code == 422 and "Nope Person" in r.json()["detail"]
    assert db_session.query(SelectionBatch).count() == 0


def test_unknown_entry_404_and_repeat_422(client, db_session, admin):
    r = client.post(URL, json={"members": [{"talent_pool_entry_id": str(uuid.uuid4())}]}, headers=admin)
    assert r.status_code == 404
    e = _entry(db_session)
    assert _create(client, admin, [e, e]).status_code == 422


def test_duplicate_membership_409_names_batch_and_withdraw_frees(client, db_session, admin):
    e1, e2 = _entry(db_session, full_name="Dup One"), _entry(db_session, full_name="Free Two")
    first = _create(client, admin, [e1]).json()
    r = _create(client, admin, [e1, e2])
    assert r.status_code == 409 and first["batch_no"] in r.json()["detail"]
    assert db_session.query(SelectionBatch).count() == 1  # nothing partial

    m = client.get(f"{URL}/membership", params={"entry_ids": [str(e1.id), str(e2.id)]}, headers=admin)
    assert m.json() == {str(e1.id): first["batch_no"]}

    w = client.post(f"{URL}/{first['id']}/withdraw", json={"reason": "wrong list"}, headers=admin)
    assert w.status_code == 200 and w.json()["status"] == "WITHDRAWN"
    assert client.get(f"{URL}/membership", params={"entry_ids": [str(e1.id)]}, headers=admin).json() == {}
    again = _create(client, admin, [e1, e2])
    assert again.status_code == 201 and again.json()["batch_no"].endswith("-002")  # withdrawn still counts
    # the withdrawn batch keeps its snapshot
    assert len(client.get(f"{URL}/{first['id']}", headers=admin).json()["members"]) == 1


def test_withdraw_without_body_and_only_when_sent(client, db_session, admin):
    b = _create(client, admin, [_entry(db_session)]).json()
    assert client.post(f"{URL}/{b['id']}/withdraw", headers=admin).status_code == 200
    assert client.post(f"{URL}/{b['id']}/withdraw", headers=admin).status_code == 409
    assert client.post(f"{URL}/{b['id']}/approval", json={"university_ref_no": "U/1", "university_approved_on": "2026-11-02"},
                       headers=admin).status_code == 409  # fmt: skip


def test_approval_rules(client, db_session, admin):
    b = _create(client, admin, [_entry(db_session)]).json()
    url = f"{URL}/{b['id']}/approval"
    assert client.post(url, json={"university_approved_on": "2026-11-02"}, headers=admin).status_code == 422
    assert client.post(url, json={"university_ref_no": "   ", "university_approved_on": "2026-11-02"}, headers=admin).status_code == 422
    assert client.post(url, json={"university_ref_no": "U/1"}, headers=admin).status_code == 422
    ok = client.post(url, json={"university_ref_no": " UNI/2026/17 ", "university_approved_on": "2026-11-02"}, headers=admin)
    assert ok.status_code == 200, ok.text
    d = ok.json()
    assert (d["status"], d["university_ref_no"], d["university_approved_on"]) == ("UNIVERSITY_APPROVED", "UNI/2026/17", "2026-11-02")
    assert d["approval_recorded_by_id"]
    # not twice, and an approved batch cannot be withdrawn
    assert client.post(url, json={"university_ref_no": "U", "university_approved_on": "2026-11-03"}, headers=admin).status_code == 409
    assert client.post(f"{URL}/{b['id']}/withdraw", headers=admin).status_code == 409


def test_numbering_is_per_year(client, db_session, admin, user_factory):
    from app.services.selection_batches import next_batch_no

    u = user_factory(UserRoleEnum.SUPER_ADMIN)
    for no in ("SB-2025-001", "SB-2025-002", "SB-2026-001"):
        db_session.add(SelectionBatch(batch_no=no, status="WITHDRAWN", created_by_id=u.id, sent_on=date(2026, 1, 1)))
    db_session.flush()
    assert next_batch_no(db_session, today=date(2025, 6, 1)) == "SB-2025-003"
    assert next_batch_no(db_session, today=date(2026, 6, 1)) == "SB-2026-002"
    assert next_batch_no(db_session, today=date(2027, 6, 1)) == "SB-2027-001"


def test_list_newest_first_with_counts_and_filter(client, db_session, admin):
    a = _create(client, admin, [_entry(db_session), _entry(db_session)]).json()
    b = _create(client, admin, [_entry(db_session)]).json()
    client.post(f"{URL}/{a['id']}/withdraw", headers=admin)
    r = client.get(URL, headers=admin).json()
    assert r["total"] == 2 and [i["batch_no"] for i in r["items"]] == [b["batch_no"], a["batch_no"]]
    assert [i["member_count"] for i in r["items"]] == [1, 2]
    only = client.get(URL, params={"status": "WITHDRAWN"}, headers=admin).json()
    assert [i["batch_no"] for i in only["items"]] == [a["batch_no"]]
    assert client.get(f"{URL}/{uuid.uuid4()}", headers=admin).status_code == 404


@pytest.mark.parametrize("role", [r for r in UserRoleEnum if r not in ALLOWED])
def test_other_roles_403(client, user_factory, db_session, role):
    campus = "SSE" if role in (UserRoleEnum.CAMPUS_HOD, UserRoleEnum.RECRUITMENT_OFFICER, UserRoleEnum.INTERVIEW_PANEL_MEMBER) else None
    h = auth_headers(client, user_factory(role, campus_code=campus))
    e = _entry(db_session)
    some = uuid.uuid4()
    assert client.get(URL, headers=h).status_code == 403
    assert _create(client, h, [e]).status_code == 403
    assert client.get(f"{URL}/membership", params={"entry_ids": [str(e.id)]}, headers=h).status_code == 403
    for method, path in (("get", ""), ("get", "/export.xlsx"), ("get", "/export.pdf"), ("post", "/withdraw"), ("post", "/approval")):
        assert getattr(client, method)(f"{URL}/{some}{path}", headers=h).status_code == 403
    assert db_session.query(SelectionBatch).count() == 0


@pytest.mark.parametrize("role", ALLOWED)
def test_allowed_roles_can_create(client, user_factory, db_session, role):
    h = auth_headers(client, user_factory(role))
    assert _create(client, h, [_entry(db_session)]).status_code == 201


def test_xlsx_and_pdf_exports(client, db_session, admin, campus_factory):
    campus = campus_factory("SSE")
    names = ["Anitha Rajan", "=Bala Formula"]
    entries = [_entry(db_session, full_name=n, dept_group="CSE") for n in names]
    b = _create(client, admin, entries, campus_id=str(campus.id)).json()

    x = client.get(f"{URL}/{b['id']}/export.xlsx", headers=admin)
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(io.BytesIO(x.content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[2][:4] == ["S.No", "Name", "Designation", "Department"]
    assert rows[3][1] == "Anitha Rajan" and rows[3][5] == "SSE"
    assert rows[4][1] == "'=Bala Formula"  # hardened, never a formula
    assert b["batch_no"] in rows[0][0]

    p = client.get(f"{URL}/{b['id']}/export.pdf", headers=admin)
    assert p.status_code == 200 and p.headers["content-type"] == "application/pdf" and p.content[:5] == b"%PDF-"
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(p.content)).pages)
    assert "Anitha Rajan" in text and "Bala" in text and b["batch_no"] in text and "Proposed salary" in text


def test_pdf_paginates_many_members(client, db_session, admin):
    entries = [_entry(db_session, full_name=f"Person Number{i:03d}") for i in range(30)]
    b = _create(client, admin, entries).json()
    p = client.get(f"{URL}/{b['id']}/export.pdf", headers=admin)
    reader = PdfReader(io.BytesIO(p.content))
    assert len(reader.pages) == 3
    text = "".join(page.extract_text() for page in reader.pages)
    assert "Person Number000" in text and "Person Number029" in text


def test_audit_rows_written_with_actor(client, db_session, user_factory):
    user = user_factory(UserRoleEnum.RECRUITMENT_COORDINATOR)
    h = auth_headers(client, user)
    b = _create(client, h, [_entry(db_session)]).json()
    client.post(f"{URL}/{b['id']}/approval", json={"university_ref_no": "U/9", "university_approved_on": "2026-11-02"}, headers=h)
    b2 = _create(client, h, [_entry(db_session)]).json()
    client.post(f"{URL}/{b2['id']}/withdraw", json={"reason": "typo"}, headers=h)
    rows = db_session.query(AuditLog).filter(AuditLog.entity_type == "SelectionBatch").all()
    assert sorted(r.action for r in rows) == [
        "SELECTION_BATCH_CREATED", "SELECTION_BATCH_CREATED",
        "SELECTION_BATCH_UNIVERSITY_APPROVED", "SELECTION_BATCH_WITHDRAWN",
    ]  # fmt: skip
    assert all(r.actor_user_id == user.id for r in rows)
    assert db_session.query(SelectionBatchMember).count() == 2
