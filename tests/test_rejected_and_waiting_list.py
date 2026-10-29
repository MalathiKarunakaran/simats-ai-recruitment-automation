"""Rejected & Waiting List (pipeline-completion brief, Phase 1).

The rules live in pipeline.transition_application_status; the page reads
/application-outcomes. Waitlisting needs INTERVIEWED and a rank, a
waitlisted application leaves only by promotion (with a reason), rejection
or withdrawal, and a rejection records its category and the stage it
happened at -- while a rejection sent the old way (free text only) still
works exactly as before.
"""

import io

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.models.audit_log import AuditLog
from app.models.enums import (
    ApplicationStatusEnum,
    HiringSlotStatusEnum,
    RejectionCategoryEnum,
    StaffRoleCategoryEnum,
    UserRoleEnum,
)
from app.models.hiring_slot import HiringSlot
from app.services import pipeline
from app.services.tracker_import import VALID_APPLICATION_STATUSES
from app.services.tracker_template import STATUSES as TEMPLATE_STATUSES

from tests.conftest import auth_headers

S = ApplicationStatusEnum


def _interviewed(db_session, application, actor):
    for target in (S.SCREENING, S.CALLED_FOR_INTERVIEW, S.INTERVIEWED):
        pipeline.transition_application_status(db_session, application=application, new_status=target, actor=actor)
    return application


def _waitlist(db_session, application, actor, rank):
    return pipeline.transition_application_status(
        db_session, application=application, new_status=S.WAITLISTED, actor=actor, waitlist_rank=rank
    )


@pytest.fixture()
def vacancy(published_vacancy_factory):
    return published_vacancy_factory(campus_code="SSE", slot_count=1)


# --- Rules ----------------------------------------------------------------


def test_waitlist_needs_interviewed_and_a_rank(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = application_factory(vacancy.job_posting, recorded_by=actor)

    with pytest.raises(HTTPException) as exc:
        _waitlist(db_session, application, actor, rank=1)
    assert exc.value.status_code == 409  # still APPLIED

    _interviewed(db_session, application, actor)
    with pytest.raises(HTTPException) as exc:
        _waitlist(db_session, application, actor, rank=None)
    assert exc.value.status_code == 422

    _waitlist(db_session, application, actor, rank=2)
    assert application.status == S.WAITLISTED
    assert application.waitlist_rank == 2


def test_waitlisting_reserves_no_slot(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, application, actor, rank=1)
    statuses = {s.status for s in db_session.query(HiringSlot).filter_by(approved_vacancy_id=vacancy.approved_vacancy.id)}
    assert statuses == {HiringSlotStatusEnum.OPEN}


@pytest.mark.parametrize("target", [S.OFFER_SENT, S.JOINED, S.INTERVIEWED])
def test_waitlisted_cannot_skip_past_selected(db_session, vacancy, application_factory, target):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, application, actor, rank=1)
    with pytest.raises(HTTPException) as exc:
        pipeline.transition_application_status(
            db_session, application=application, new_status=target, actor=actor, reason="x"
        )
    assert exc.value.status_code == 409


def test_promote_needs_a_reason_and_is_logged(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, application, actor, rank=1)

    with pytest.raises(HTTPException) as exc:
        pipeline.transition_application_status(db_session, application=application, new_status=S.SELECTED, actor=actor)
    assert exc.value.status_code == 422

    pipeline.transition_application_status(
        db_session, application=application, new_status=S.SELECTED, actor=actor, reason="First choice declined"
    )
    assert application.status == S.SELECTED
    slot = db_session.query(HiringSlot).filter_by(reserved_application_id=application.id).one()
    assert slot.status == HiringSlotStatusEnum.RESERVED

    row = (
        db_session.query(AuditLog)
        .filter_by(entity_id=application.id, action="APPLICATION_STATUS_TRANSITION")
        .order_by(AuditLog.created_at.desc(), AuditLog.id)
        .all()
    )
    promoted = next(r for r in row if r.after_state["status"] == "SELECTED")
    assert promoted.before_state["status"] == "WAITLISTED"
    assert promoted.after_state["reason"] == "First choice declined"
    assert promoted.actor_user_id == actor.id


def test_promote_is_refused_while_the_slot_is_taken_then_allowed_once_freed(
    db_session, vacancy, application_factory
):
    actor = vacancy.hr_admin
    first = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    pipeline.transition_application_status(db_session, application=first, new_status=S.SELECTED, actor=actor)
    reserve = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, reserve, actor, rank=1)

    # A refused request rolls back; the savepoint stands in for that here.
    savepoint = db_session.begin_nested()
    with pytest.raises(HTTPException) as exc:
        pipeline.transition_application_status(
            db_session, application=reserve, new_status=S.SELECTED, actor=actor, reason="Promote"
        )
    assert exc.value.status_code == 409  # the only slot is reserved by `first`
    savepoint.rollback()
    assert reserve.status == S.WAITLISTED

    pipeline.transition_application_status(
        db_session, application=first, new_status=S.WITHDRAWN, actor=actor, reason="Took another offer"
    )
    pipeline.transition_application_status(
        db_session, application=reserve, new_status=S.SELECTED, actor=actor, reason="Promote"
    )
    assert reserve.status == S.SELECTED


def test_reject_records_category_and_stage(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    pipeline.transition_application_status(
        db_session,
        application=application,
        new_status=S.REJECTED,
        actor=actor,
        reason="Panel felt depth was lacking",
        rejection_category=RejectionCategoryEnum.PANEL_NOT_RECOMMENDED,
    )
    assert application.rejection_category == RejectionCategoryEnum.PANEL_NOT_RECOMMENDED
    assert application.rejected_at_stage == S.INTERVIEWED
    assert application.rejection_reason == "Panel felt depth was lacking"


def test_reject_with_free_text_only_still_works(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = application_factory(vacancy.job_posting, recorded_by=actor)
    pipeline.transition_application_status(
        db_session, application=application, new_status=S.REJECTED, actor=actor, reason="Not a fit"
    )
    assert application.status == S.REJECTED
    assert application.rejection_category is None
    assert application.rejected_at_stage == S.APPLIED


def test_a_waitlisted_application_can_be_rejected(db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, application, actor, rank=1)
    pipeline.transition_application_status(
        db_session,
        application=application,
        new_status=S.REJECTED,
        actor=actor,
        reason="Vacancy filled",
        rejection_category=RejectionCategoryEnum.OTHER,
    )
    assert application.rejected_at_stage == S.WAITLISTED


def test_tracker_import_and_template_still_refuse_waitlisted():
    assert "WAITLISTED" not in VALID_APPLICATION_STATUSES
    assert "WAITLISTED" not in TEMPLATE_STATUSES


# --- API ------------------------------------------------------------------


def test_status_endpoint_accepts_rank_and_category(client, db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    db_session.commit()
    headers = auth_headers(client, actor)

    response = client.patch(
        f"/api/v1/applications/{application.id}/status",
        headers=headers,
        json={"status": "WAITLISTED", "waitlist_rank": 3},
    )
    assert response.status_code == 200, response.text
    assert response.json()["waitlist_rank"] == 3

    response = client.patch(
        f"/api/v1/applications/{application.id}/status",
        headers=headers,
        json={"status": "REJECTED", "reason": "Filled", "rejection_category": "SALARY_MISMATCH"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["rejection_category"] == "SALARY_MISMATCH"
    assert body["rejected_at_stage"] == "WAITLISTED"


def test_rank_below_one_is_refused(client, db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    application = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    db_session.commit()
    response = client.patch(
        f"/api/v1/applications/{application.id}/status",
        headers=auth_headers(client, actor),
        json={"status": "WAITLISTED", "waitlist_rank": 0},
    )
    assert response.status_code == 422


def test_waiting_list_is_ranked_and_campus_scoped(
    client, db_session, published_vacancy_factory, application_factory, user_factory
):
    sse = published_vacancy_factory(campus_code="SSE", slot_count=1)
    sclas = published_vacancy_factory(campus_code="SCLAS", slot_count=1)
    actor = sse.hr_admin
    second = _interviewed(db_session, application_factory(sse.job_posting, recorded_by=actor), actor)
    first = _interviewed(db_session, application_factory(sse.job_posting, recorded_by=actor), actor)
    other = _interviewed(db_session, application_factory(sclas.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, second, actor, rank=2)
    _waitlist(db_session, first, actor, rank=1)
    _waitlist(db_session, other, actor, rank=1)
    db_session.commit()

    body = client.get("/api/v1/application-outcomes/waiting-list", headers=auth_headers(client, actor)).json()
    sse_rows = [r for r in body["items"] if r["job_posting_id"] == str(sse.job_posting.id)]
    assert [r["application_id"] for r in sse_rows] == [str(first.id), str(second.id)]
    assert body["total"] == 3

    officer = user_factory(UserRoleEnum.RECRUITMENT_OFFICER, campus_code="SSE")
    db_session.commit()
    body = client.get("/api/v1/application-outcomes/waiting-list", headers=auth_headers(client, officer)).json()
    assert {r["campus_code"] for r in body["items"]} == {"SSE"}
    assert body["total"] == 2


def test_rejected_filters_and_export(client, db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    screened = application_factory(vacancy.job_posting, recorded_by=actor)
    pipeline.transition_application_status(db_session, application=screened, new_status=S.SCREENING, actor=actor)
    pipeline.transition_application_status(
        db_session,
        application=screened,
        new_status=S.REJECTED,
        actor=actor,
        reason="No PG",
        rejection_category=RejectionCategoryEnum.NOT_ELIGIBLE,
    )
    interviewed = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    pipeline.transition_application_status(
        db_session,
        application=interviewed,
        new_status=S.REJECTED,
        actor=actor,
        reason="Weak demo",
        rejection_category=RejectionCategoryEnum.PANEL_NOT_RECOMMENDED,
    )
    db_session.commit()
    headers = auth_headers(client, actor)

    body = client.get(
        "/api/v1/application-outcomes/rejected", headers=headers, params={"rejection_category": "NOT_ELIGIBLE"}
    ).json()
    assert [r["application_id"] for r in body["items"]] == [str(screened.id)]

    body = client.get(
        "/api/v1/application-outcomes/rejected", headers=headers, params={"rejected_at_stage": "INTERVIEWED"}
    ).json()
    assert [r["rejection_remarks"] for r in body["items"]] == ["Weak demo"]

    response = client.get("/api/v1/application-outcomes/rejected/export", headers=headers)
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(i for i, r in enumerate(rows) if r and r[0] == "Candidate")
    data = rows[header_index + 1 :]
    assert len(data) == 2
    assert {r[8] for r in data} == {"NOT_ELIGIBLE", "PANEL_NOT_RECOMMENDED"}


def test_category_counts_ignore_the_selected_category(
    client, db_session, published_vacancy_factory, application_factory
):
    teaching = published_vacancy_factory(campus_code="SSE", slot_count=1)
    nts = published_vacancy_factory(
        campus_code="SSE", slot_count=1, role_category=StaffRoleCategoryEnum.NON_TEACHING
    )
    actor = teaching.hr_admin
    for vacancy_ in (teaching, teaching, nts):
        application = application_factory(vacancy_.job_posting, recorded_by=actor)
        pipeline.transition_application_status(
            db_session, application=application, new_status=S.REJECTED, actor=actor, reason="x"
        )
    db_session.commit()

    body = client.get(
        "/api/v1/application-outcomes/rejected",
        headers=auth_headers(client, actor),
        params={"role_category": "NON_TEACHING"},
    ).json()
    assert body["total"] == 1
    assert body["category_counts"] == {"TEACHING": 2, "NON_TEACHING": 1, "HOUSEKEEPING": 0, "ALL": 3}


def test_candidate_cannot_read_outcomes(client, db_session, user_factory):
    candidate_user = user_factory(UserRoleEnum.CANDIDATE)
    db_session.commit()
    response = client.get("/api/v1/application-outcomes/rejected", headers=auth_headers(client, candidate_user))
    assert response.status_code == 403


# --- Dashboard ------------------------------------------------------------


def test_dashboard_counts_waitlisted_apart_from_rejected(client, db_session, vacancy, application_factory):
    actor = vacancy.hr_admin
    waiting = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    _waitlist(db_session, waiting, actor, rank=1)
    categorised = application_factory(vacancy.job_posting, recorded_by=actor)
    pipeline.transition_application_status(
        db_session,
        application=categorised,
        new_status=S.REJECTED,
        actor=actor,
        reason="No PG",
        rejection_category=RejectionCategoryEnum.NOT_ELIGIBLE,
    )
    legacy = application_factory(vacancy.job_posting, recorded_by=actor)
    pipeline.transition_application_status(
        db_session, application=legacy, new_status=S.REJECTED, actor=actor, reason="Not a fit"
    )
    db_session.commit()

    body = client.get("/api/v1/dashboard/kpis", headers=auth_headers(client, actor)).json()
    assert body["waitlisted_count"] == 1
    assert body["rejected_count"] == 2  # the waitlisted one is not in here
    assert {(r["category"], r["count"]) for r in body["rejected_by_reason"]} == {("NOT_ELIGIBLE", 1), (None, 1)}
    funnel = {r["stage"]: r["count"] for r in body["application_pipeline_funnel"]}
    assert funnel["Waitlisted"] == 1
    assert funnel["Rejected"] == 2


def test_waiting_list_rows_say_whether_a_post_is_free(client, db_session, vacancy, application_factory, candidate_factory):
    actor = vacancy.hr_admin
    holder = _interviewed(
        db_session,
        application_factory(vacancy.job_posting, recorded_by=actor, candidate=candidate_factory(full_name="Priya Raman")),
        actor,
    )
    pipeline.transition_application_status(db_session, application=holder, new_status=S.SELECTED, actor=actor)
    reserve = _interviewed(db_session, application_factory(vacancy.job_posting, recorded_by=actor), actor)
    reserve.panel_remarks = "Strong research record"
    _waitlist(db_session, reserve, actor, rank=1)
    db_session.commit()
    headers = auth_headers(client, actor)

    row = client.get("/api/v1/application-outcomes/waiting-list", headers=headers).json()["items"][0]
    assert row["vacancy_total_positions"] == 1
    assert row["vacancy_open_slots"] == 0
    assert row["vacancy_slot_holders"] == ["Priya Raman"]
    assert row["panel_remarks"] == "Strong research record"

    pipeline.transition_application_status(
        db_session, application=holder, new_status=S.WITHDRAWN, actor=actor, reason="Took another offer"
    )
    db_session.commit()
    row = client.get("/api/v1/application-outcomes/waiting-list", headers=headers).json()["items"][0]
    assert row["vacancy_open_slots"] == 1
    assert row["vacancy_slot_holders"] == []
