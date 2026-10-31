"""New Joinees + Pro App Order sync (docs/specs/joinee-order-sync.md step 2)."""

import io
import logging
from datetime import date

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.models.audit_log import AuditLog
from app.models.enums import TalentPoolSourceTabEnum, UserRoleEnum
from app.models.sheet_sync import ProvisionalOrder, SheetJoinee, SheetSyncRun
from app.models.talent_pool_entry import TalentPoolEntry

from tests.conftest import auth_headers

SECRET = "test-talent-pool-sync-secret"
HEADERS = {"X-Talent-Pool-Sync-Secret": SECRET}

J_HEADERS = [
    "S. No+", "BIO ID", "New BIoID", "Staff ID", "App No", "Name", "Gender", "Employee Type", "College", "Status",
    "Joined On", "Interview Date", "Provisional Appointment Order Given Date", "Designation", "Department",
    "Dept New", "Reporting to...", "Relieving Order Submitted", "Original Submitted", "Appointment Order",
    "send date ", "Received date", "Mobile No.", "RBL Account number", "Salary fixed", "Allowance Details",
    "Current Work",
]  # fmt: skip
O_HEADERS = [
    # The live tab's own header row (30-Oct-2026), in its order, plus a
    # column that must be ignored.
    "S.No", "Order Date", "BioID", "EMPID", "Name", "Provision Order", "Probation Period", "ID Card", "Food App",
    "Arms", "360 App", "File Submitted to ADFA", "Room Allotted", "After Probation  Received AO ", "Salary fixed",
]  # fmt: skip


@pytest.fixture(autouse=True)
def _secret(monkeypatch):
    monkeypatch.setattr(settings, "TALENT_POOL_SYNC_SECRET", SECRET)


def jrow(n, headers=J_HEADERS, **vals):
    base = {"BIO ID": f"B{n}", "Name": f"Person {n}", "Employee Type": "TS", "College": "SSE", "Status": "New",
            "Joined On": "05-06-2025"}  # fmt: skip
    base.update(vals)
    return {"row_number": n + 2, "cells": [base.get(h.strip(), "") if h not in base else base[h] for h in headers]}


def orow(n, **vals):
    base = {"BioID": f"B{n}", "EMPID": f"E{n}", "Name": f"Person {n}", "Order Date": "01-02-2025",
            "Probation Period": "2/28/2026", "Provision Order": "Issued", "After Probation  Received AO ": "Received"}  # fmt: skip
    base.update(vals)
    return {"row_number": n + 1, "cells": [base.get(h, "") for h in O_HEADERS]}


def post(client, source, rows, headers=None, dry_run=False, secret=SECRET):
    payload = {"headers": headers or (J_HEADERS if source == "NEW_JOINEES" else O_HEADERS), "rows": rows, "dry_run": dry_run}
    h = {"X-Talent-Pool-Sync-Secret": secret} if secret is not None else {}
    return client.post(f"/api/v1/sheet-sync/{source}", json=payload, headers=h)


def joinees(db):
    db.expire_all()
    return db.query(SheetJoinee).order_by(SheetJoinee.sheet_row_number).all()


def orders(db):
    db.expire_all()
    return db.query(ProvisionalOrder).order_by(ProvisionalOrder.sheet_row_number).all()


# --- Gate -----------------------------------------------------------------


def test_unset_secret_503(client, monkeypatch):
    monkeypatch.setattr(settings, "TALENT_POOL_SYNC_SECRET", "")
    assert post(client, "NEW_JOINEES", [jrow(1)]).status_code == 503


@pytest.mark.parametrize("secret", [None, "wrong"])
def test_bad_secret_401(client, db_session, secret):
    assert post(client, "NEW_JOINEES", [jrow(1)], secret=secret).status_code == 401
    assert joinees(db_session) == []


def test_missing_name_header_422(client, db_session):
    h = [x for x in J_HEADERS if x != "Name"]
    r = post(client, "NEW_JOINEES", [{"row_number": 3, "cells": ["x"] * len(h)}], headers=h)
    assert r.status_code == 422
    assert joinees(db_session) == []


def test_unknown_source_422(client):
    assert post(client, "OTHER", []).status_code == 422


# --- Lifecycle ------------------------------------------------------------


def test_insert_update_unchanged_deactivate_reactivate(client, db_session):
    r = post(client, "NEW_JOINEES", [jrow(1), jrow(2)])
    assert r.status_code == 200, r.text
    assert r.json()["counts"]["inserted"] == 2
    assert r.json()["counts"]["not_in_employees"] == 2

    r = post(client, "NEW_JOINEES", [jrow(1), jrow(2)])
    c = r.json()["counts"]
    assert (c["inserted"], c["updated"], c["unchanged"]) == (0, 0, 2)

    r = post(client, "NEW_JOINEES", [jrow(1, Designation="Professor"), jrow(2)])
    assert r.json()["counts"]["updated"] == 1
    assert joinees(db_session)[0].designation == "Professor"

    r = post(client, "NEW_JOINEES", [jrow(1, Designation="Professor")])
    assert r.json()["counts"]["deactivated"] == 1
    assert [j.is_active for j in joinees(db_session)] == [True, False]

    r = post(client, "NEW_JOINEES", [jrow(1, Designation="Professor"), jrow(2)])
    assert r.json()["counts"]["reactivated"] == 1
    assert all(j.is_active for j in joinees(db_session))


def test_moved_row_is_unchanged(client, db_session):
    post(client, "NEW_JOINEES", [jrow(1), jrow(2)])
    moved = jrow(2)
    moved["row_number"] = 3
    first = jrow(1)
    first["row_number"] = 4
    r = post(client, "NEW_JOINEES", [moved, first])
    assert r.json()["counts"]["unchanged"] == 2


def test_dry_run_changes_nothing_but_records_run(client, db_session):
    r = post(client, "NEW_JOINEES", [jrow(1)], dry_run=True)
    assert r.status_code == 200 and r.json()["dry_run"] is True
    assert r.json()["counts"]["inserted"] == 1
    assert joinees(db_session) == []
    run = db_session.query(SheetSyncRun).one()
    assert run.dry_run is True and run.counts["inserted"] == 1
    # a real run afterwards still inserts
    assert post(client, "NEW_JOINEES", [jrow(1)]).json()["counts"]["inserted"] == 1


# --- Keys -------------------------------------------------------------------


def test_key_fallback_chain_and_repeats(client, db_session):
    rows = [
        jrow(1, **{"BIO ID": "", "New BIoID": "NB1"}),
        jrow(2, **{"BIO ID": "", "Staff ID": "S2"}),
        jrow(3, **{"BIO ID": "", "Name": "No Ids", "Joined On": "01-01-2024"}),
        jrow(4, **{"BIO ID": "DUP"}),
        jrow(5, **{"BIO ID": "DUP", "Status": "Rejoin"}),
    ]
    r = post(client, "NEW_JOINEES", rows)
    c = r.json()["counts"]
    assert c["inserted"] == 5 and c["duplicates"] == 1
    assert any("key repeats" in f["reason"] for f in r.json()["flagged"])
    # idempotent with repeats
    assert post(client, "NEW_JOINEES", rows).json()["counts"]["unchanged"] == 5


def test_no_key_row_skipped_and_flagged(client, db_session):
    r = post(client, "NEW_JOINEES", [jrow(1, **{"BIO ID": "", "Name": "", "Designation": "X"})])
    assert r.json()["counts"]["skipped"] == 1
    assert any("no key" in f["reason"] for f in r.json()["flagged"])
    assert joinees(db_session) == []


# --- Normalisation ----------------------------------------------------------


@pytest.mark.parametrize(
    "cell,expected,flagged",
    [("TS", "TEACHING", False), ("NTS", "NON_TEACHING", False), ("HK", "HOUSEKEEPING", False),
     ("JRF", None, True), ("Weird", None, True)],
)  # fmt: skip
def test_category(client, db_session, cell, expected, flagged):
    r = post(client, "NEW_JOINEES", [jrow(1, **{"Employee Type": cell})])
    j = joinees(db_session)[0]
    assert j.category == expected and j.employee_type_raw == cell
    assert any("Employee Type" in f["reason"] for f in r.json()["flagged"]) is flagged


@pytest.mark.parametrize(
    "cell,expected",
    [("SSE", "SSE"), ("SCLAS", "SCLAS"), ("SCALS", "SCLAS"), ("SCAD", "SCAD"), ("SSPE", "SSPE"),
     ("STUDIO", "STUDIO"), ("SPIER", "SPIER"), ("SHIFT", "SHIFT"), ("SHoTs", None), ("SNAPS", None), ("Zed", None)],
)  # fmt: skip
def test_campus(client, db_session, campus_factory, cell, expected):
    if expected:
        campus = campus_factory(expected)
    r = post(client, "NEW_JOINEES", [jrow(1, College=cell)])
    j = joinees(db_session)[0]
    assert j.campus_code == expected and j.college_raw == cell
    assert (j.campus_id == campus.id) if expected else j.campus_id is None
    assert any("College" in f["reason"] for f in r.json()["flagged"]) is (expected is None)


@pytest.mark.parametrize(
    "cell,state,flagged",
    [("New", "ACTIVE", False), ("Long Leave", "ACTIVE", True), ("Left", "LEFT", False), ("LEFT", "LEFT", False),
     ("Resigned ", "LEFT", False), ("EXPIRED", "LEFT", False), ("Rejoin", "REJOINED", False),
     ("Rejoined", "REJOINED", False), ("Rejoned", "REJOINED", False), ("REJOIN", "REJOINED", False),
     ("Old", None, True)],
)  # fmt: skip
def test_employment_state(client, db_session, cell, state, flagged):
    r = post(client, "NEW_JOINEES", [jrow(1, Status=cell)])
    assert joinees(db_session)[0].employment_state == state
    assert any("Status" in f["reason"] or "Long Leave" in f["reason"] for f in r.json()["flagged"]) is flagged


def test_joiner_dates(client, db_session):
    r = post(client, "NEW_JOINEES", [jrow(1, **{"Joined On": "31-12-2024", "Interview Date": "12/31/2024",
                                               "Received date": "05-06-2025", "send date ": "04-06-2025"})])  # fmt: skip
    j = joinees(db_session)[0]
    assert j.joined_on == date(2024, 12, 31)
    assert j.interview_on is None and j.interview_on_raw == "12/31/2024"
    assert (j.send_on, j.received_on) == (date(2025, 6, 4), date(2025, 6, 5))
    assert any("Interview Date" in f["reason"] for f in r.json()["flagged"])


# --- Pro App Order ------------------------------------------------------------


def test_order_normalisation_and_dates(client, db_session):
    r = post(client, "PRO_APP_ORDER", [orow(1)])
    assert r.status_code == 200, r.text
    o = orders(db_session)[0]
    assert o.order_on == date(2025, 2, 1)  # DD-MM-YYYY
    assert o.probation_end_on == date(2026, 2, 28)  # M/D/YYYY
    assert (o.order_status, o.outcome) == ("ISSUED", "PERMANENT_ORDER_RECEIVED")
    assert r.json()["flagged"] == []


@pytest.mark.parametrize(
    "status,expected", [("Issued", "ISSUED"), ("Generated", "ISSUED"), ("Not Issued", "PENDING"), ("", "PENDING")]
)
def test_order_status(client, db_session, status, expected):
    post(client, "PRO_APP_ORDER", [orow(1, **{"Provision Order": status})])
    assert orders(db_session)[0].order_status == expected


@pytest.mark.parametrize(
    "cell,expected,flagged",
    [("Received ", "PERMANENT_ORDER_RECEIVED", False), ("Extend Probation", "PROBATION_EXTENDED", False),
     ("Not Satisfied", "NOT_RECOMMENDED", False), ("", None, False), ("Left", None, True), ("Rejoin", None, True),
     ("TRA", None, True), ("Family Way", None, True), ("No Submit PG", None, True), ("Long Leave", None, True)],
)  # fmt: skip
def test_order_outcome(client, db_session, cell, expected, flagged):
    r = post(client, "PRO_APP_ORDER", [orow(1, **{"After Probation  Received AO ": cell})])
    assert orders(db_session)[0].outcome == expected
    assert any("Outcome" in f["reason"] for f in r.json()["flagged"]) is flagged


def test_room_left_flagged_and_bad_dates(client, db_session):
    r = post(client, "PRO_APP_ORDER", [orow(1, **{"Room Allotted": "Left", "Order Date": "2/1/2025",
                                                 "Probation Period": "28-02-2026"})])  # fmt: skip
    o = orders(db_session)[0]
    assert o.room_allotted == "Left" and o.order_on is None and o.probation_end_on is None
    reasons = " ".join(f["reason"] for f in r.json()["flagged"])
    assert "Room" in reasons and "Order Date" in reasons and "Probation Period" in reasons


def test_order_links_to_joiner_and_idempotent(client, db_session):
    post(client, "NEW_JOINEES", [jrow(1)])
    r = post(client, "PRO_APP_ORDER", [orow(1), orow(2)])
    assert r.json()["counts"]["linked_joinee"] == 1 and r.json()["counts"]["unlinked"] == 1
    by_bio = {o.bio_id: o for o in orders(db_session)}
    assert by_bio["B1"].sheet_joinee_id == joinees(db_session)[0].id
    assert by_bio["B2"].sheet_joinee_id is None
    c = post(client, "PRO_APP_ORDER", [orow(1), orow(2)]).json()["counts"]
    assert (c["inserted"], c["updated"], c["unchanged"]) == (0, 0, 2)
    assert post(client, "PRO_APP_ORDER", [orow(1)]).json()["counts"]["deactivated"] == 1


# --- Talent pool link -------------------------------------------------------


def test_talent_pool_link_by_phone_last_10(client, db_session):
    keyed = TalentPoolEntry(source_tab=TalentPoolSourceTabEnum.TEACHING_FINAL, full_name="Pool A", phone="+91 98765-43210",
                            sheet_status="Pending", row_key="k1")
    unkeyed = TalentPoolEntry(source_tab=TalentPoolSourceTabEnum.TEACHING_FINAL, full_name="Pool B", phone="9111122223",
                              sheet_status="Pending")
    db_session.add_all([keyed, unkeyed])
    db_session.flush()
    post(client, "NEW_JOINEES", [jrow(1, **{"Mobile No.": "098765 43210"}), jrow(2, **{"Mobile No.": "9111122223"})])
    j = joinees(db_session)
    assert j[0].talent_pool_entry_id == keyed.id
    assert j[1].talent_pool_entry_id is None  # hand-entered rows are never linked


# --- Privacy ------------------------------------------------------------------


def test_excluded_columns_and_mobile_never_stored(client, db_session, user_factory, caplog):
    sentinels = {
        "RBL Account number": "RBLSENTINEL7781", "Salary fixed": "SALSENTINEL9942",
        "Allowance Details": "ALLOWSENTINEL3310", "Gender": "GENDERSENTINEL5521",
        "Current Work": "WORKSENTINEL6643", "S. No+": "SNOSENTINEL1188", "Mobile No.": "9000011111",
    }  # fmt: skip
    caplog.set_level(logging.DEBUG)
    row = jrow(1, **sentinels, **{"Employee Type": "JRF", "College": "SHoTs", "Status": "Old"})
    r = post(client, "NEW_JOINEES", [row])
    assert r.status_code == 200, r.text
    # the same payload on a dry run, and one that is rejected, leak nothing either
    post(client, "NEW_JOINEES", [row], dry_run=True)
    bad_headers = [h for h in J_HEADERS if h != "Name"]
    bad = post(client, "NEW_JOINEES", [{"row_number": 3, "cells": list(sentinels.values())}], headers=bad_headers)
    assert bad.status_code == 422

    admin = user_factory(UserRoleEnum.SUPER_ADMIN)
    run_id = r.json()["run_id"]
    x = client.get(f"/api/v1/sheet-sync/runs/{run_id}/flagged.xlsx", headers=auth_headers(client, admin))
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.max_row > 1
    runs = client.get("/api/v1/sheet-sync/runs", headers=auth_headers(client, admin))
    assert runs.status_code == 200

    db_session.expire_all()
    blobs = [
        repr([(c.name, getattr(j, c.name)) for c in SheetJoinee.__table__.columns]) for j in db_session.query(SheetJoinee)
    ]
    blobs += [repr((s.counts, s.flagged, s.missing_headers)) for s in db_session.query(SheetSyncRun)]
    blobs += [repr((a.before_state, a.after_state)) for a in db_session.query(AuditLog)]
    blobs += [r.text, bad.text, runs.text, repr(x.content), caplog.text]
    # raw xlsx is zipped, so also check every cell text
    blobs.append(repr([[c.value for c in row_] for row_ in ws.iter_rows()]))
    haystack = "\n".join(blobs)
    for label, value in sentinels.items():
        assert value not in haystack, label
    for h in ("RBL Account number", "Salary fixed", "Allowance Details"):
        assert h not in haystack


def test_runs_endpoints_are_super_admin_only(client, user_factory, db_session):
    hr = user_factory(UserRoleEnum.HR_ADMIN)
    assert client.get("/api/v1/sheet-sync/runs", headers=auth_headers(client, hr)).status_code == 403
    post(client, "NEW_JOINEES", [jrow(1)])
    run = db_session.query(SheetSyncRun).one()
    assert client.get(f"/api/v1/sheet-sync/runs/{run.id}/flagged.xlsx", headers=auth_headers(client, hr)).status_code == 403


# --- S5 New Departments -------------------------------------------------------

from app.models.sheet_sync import DepartmentFaculty, NewDepartment  # noqa: E402

D_HEADERS = ["#", "Approved Department", "Notes", "Count", "Additional Count"]
F_HEADERS = ["S.No", "Bio ID", "Name", "Dept", "Remark", "Status", "Top 25", "Vacancy", "Active"]


def drow(n, **vals):
    base = {"Approved Department": f"Dept {n}", "Notes": "", "Count": "5", "Additional Count": "2"}
    base.update(vals)
    return {"row_number": n + 1, "cells": [base.get(h, "") for h in D_HEADERS]}


def frow(n, **vals):
    base = {"Bio ID": f"F{n}", "Name": f"Faculty {n}", "Dept": "Dept 1", "Remark": "Member"}
    base.update(vals)
    return {"row_number": n + 1, "cells": [base.get(h, "") for h in F_HEADERS]}


def post5(client, source, rows, dry_run=False):
    headers = D_HEADERS if source == "NEW_DEPARTMENTS" else F_HEADERS
    return post(client, source, rows, headers=headers, dry_run=dry_run)


def depts(db):
    db.expire_all()
    return db.query(NewDepartment).order_by(NewDepartment.sheet_row_number).all()


def faculty(db):
    db.expire_all()
    return db.query(DepartmentFaculty).order_by(DepartmentFaculty.sheet_row_number).all()


def test_departments_counts_notes_and_flags(client, db_session):
    r = post5(client, "NEW_DEPARTMENTS", [drow(1, Notes="Top Dept"), drow(2, Count="five", **{"Additional Count": ""}),
                                          drow(3, **{"Approved Department": ""}, Count="4")])  # fmt: skip
    assert r.status_code == 200, r.text
    c = r.json()["counts"]
    assert c["inserted"] == 2 and c["skipped"] == 1
    d1, d2 = depts(db_session)
    assert (d1.name, d1.notes, d1.count, d1.additional_count) == ("Dept 1", "Top Dept", 5, 2)
    assert d2.count is None and d2.count_raw == "five" and d2.additional_count is None
    reasons = [f["reason"] for f in r.json()["flagged"]]
    assert any("Count is not a whole number" in x for x in reasons) and any("no key" in x for x in reasons)


def test_departments_idempotent_deactivate_and_dry_run(client, db_session):
    rows = [drow(1), drow(2)]
    assert post5(client, "NEW_DEPARTMENTS", rows, dry_run=True).json()["counts"]["inserted"] == 2
    assert depts(db_session) == []
    assert post5(client, "NEW_DEPARTMENTS", rows).json()["counts"]["inserted"] == 2
    c = post5(client, "NEW_DEPARTMENTS", rows).json()["counts"]
    assert (c["inserted"], c["updated"], c["unchanged"]) == (0, 0, 2)
    assert post5(client, "NEW_DEPARTMENTS", [drow(1)]).json()["counts"]["deactivated"] == 1
    assert [d.is_active for d in depts(db_session)] == [True, False]
    assert db_session.query(SheetSyncRun).filter_by(dry_run=True, source="NEW_DEPARTMENTS").count() == 1


def test_missing_key_header_422_for_both(client):
    assert post(client, "NEW_DEPARTMENTS", [], headers=["Notes"]).status_code == 422
    assert post(client, "DEPARTMENT_FACULTY", [], headers=["Name"]).status_code == 422


@pytest.mark.parametrize(
    "cell,role,flagged",
    [("HOD", "HOD", False), ("HoD", "HOD", False), ("Member", "MEMBER", False), ("mEmber", "MEMBER", False),
     ("Dean", None, True), ("", None, True)],
)  # fmt: skip
def test_faculty_remark(client, db_session, cell, role, flagged):
    post5(client, "NEW_DEPARTMENTS", [drow(1)])
    r = post5(client, "DEPARTMENT_FACULTY", [frow(1, Remark=cell)])
    f = faculty(db_session)[0]
    assert f.role == role and f.remark_raw == (cell or None)
    assert any("Remark" in x["reason"] for x in r.json()["flagged"]) is flagged


@pytest.mark.parametrize(
    "cell,top,state,note",
    [("Top Dept", True, None, None), ("LEFT", False, "LEFT", None), (";LEFT", False, "LEFT", None),
     ("RESIGNED", False, "LEFT", None), ("LEAVE", False, "ON_LEAVE", None), ("MATERNITY", False, "ON_LEAVE", None),
     ("CSE", False, None, "CSE"), ("ECE", False, None, "ECE"), ("MATHS", False, None, "MATHS"),
     ("", False, None, None)],
)  # fmt: skip
def test_faculty_top25(client, db_session, cell, top, state, note):
    post5(client, "DEPARTMENT_FACULTY", [frow(1, **{"Top 25": cell})])
    f = faculty(db_session)[0]
    assert (f.top_dept, f.employment_state, f.parent_dept_note) == (top, state, note)
    assert f.top25_raw == (cell or None)


def test_faculty_status_additional(client, db_session):
    r = post5(client, "DEPARTMENT_FACULTY", [frow(1, Status="Addtional Faculty"), frow(2), frow(3, Status="Weird")])
    f = faculty(db_session)
    assert [x.is_additional for x in f] == [True, False, False]
    assert any("Status" in x["reason"] and x["key"] == "F3" for x in r.json()["flagged"])
    assert not any("Status" in x["reason"] and x["key"] != "F3" for x in r.json()["flagged"])


def test_faculty_dept_exact_fold_only_no_fuzzy(client, db_session):
    post5(client, "NEW_DEPARTMENTS", [drow(1, **{"Approved Department": "Additive Manufacturing Engineering"}),
                                      drow(2, **{"Approved Department": "Civil  Engineering"})])  # fmt: skip
    r = post5(client, "DEPARTMENT_FACULTY", [
        frow(1, Dept="additive manufacturing  engineering "),
        frow(2, Dept="Additive Manufacuring Engineering"),  # typo: never guessed
        frow(3, Dept="civil engineering"),
        frow(4, Dept=""),
    ])  # fmt: skip
    f = faculty(db_session)
    d = {x.name: x.id for x in depts(db_session)}
    assert f[0].new_department_id == d["Additive Manufacturing Engineering"]
    assert f[1].new_department_id is None and f[1].dept_raw == "Additive Manufacuring Engineering"
    assert f[2].new_department_id == d["Civil Engineering"]
    assert f[3].new_department_id is None
    flagged = {x["key"]: x["reason"] for x in r.json()["flagged"] if "Dept" in x["reason"]}
    assert set(flagged) == {"F2", "F4"} and "blank" in flagged["F4"]
    c = r.json()["counts"]
    assert (c["linked_department"], c["unlinked_department"]) == (2, 2)


def test_faculty_links_joiner_and_lifecycle(client, db_session):
    post(client, "NEW_JOINEES", [jrow(1, **{"BIO ID": "F1"})])
    post5(client, "NEW_DEPARTMENTS", [drow(1)])
    rows = [frow(1), frow(2)]
    r = post5(client, "DEPARTMENT_FACULTY", rows)
    assert r.json()["counts"]["linked_joinee"] == 1
    f = faculty(db_session)
    assert f[0].sheet_joinee_id == joinees(db_session)[0].id and f[1].sheet_joinee_id is None
    c = post5(client, "DEPARTMENT_FACULTY", rows).json()["counts"]
    assert (c["inserted"], c["updated"], c["unchanged"]) == (0, 0, 2)
    assert post5(client, "DEPARTMENT_FACULTY", [frow(1, Remark="HOD")]).json()["counts"]["updated"] == 1
    assert [x.is_active for x in faculty(db_session)] == [True, False]
    assert post5(client, "DEPARTMENT_FACULTY", [frow(1, Remark="HOD"), frow(2)]).json()["counts"]["reactivated"] == 1
    assert all(x.is_active for x in faculty(db_session))


def test_faculty_dry_run_changes_nothing_and_runs_list_has_new_sources(client, db_session, user_factory):
    r = post5(client, "DEPARTMENT_FACULTY", [frow(1)], dry_run=True)
    assert r.json()["dry_run"] is True and r.json()["counts"]["inserted"] == 1
    assert faculty(db_session) == []
    admin = user_factory(UserRoleEnum.SUPER_ADMIN)
    runs = client.get("/api/v1/sheet-sync/runs", headers=auth_headers(client, admin)).json()
    by_source = {x["source"]: x for x in runs}
    assert {"NEW_DEPARTMENTS", "DEPARTMENT_FACULTY"} <= set(by_source)
    assert by_source["DEPARTMENT_FACULTY"]["last_dry_run"]["counts"]["inserted"] == 1
    assert by_source["NEW_DEPARTMENTS"]["last_run"] is None


def test_s5_secret_gate(client, db_session):
    assert post5(client, "DEPARTMENT_FACULTY", [frow(1)]).status_code == 200
    assert post(client, "NEW_DEPARTMENTS", [drow(1)], headers=D_HEADERS, secret="wrong").status_code == 401
